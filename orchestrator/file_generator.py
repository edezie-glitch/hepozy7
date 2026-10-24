<<<<<<< HEAD
import csv as csv_module
import io
import os
import re
import time
import uuid
from xml.sax.saxutils import escape as _xml_escape

# CHANGED: absolute path next to this file. A relative name depends on which folder the server
# was started from, so files could land in one place and be served from another.
GENERATED_FILES_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "generated_files")
os.makedirs(GENERATED_FILES_DIR, exist_ok=True)

# ADDED: False = the file holds ONLY the content. The short title is used for the file name only.
# Set to True to print the title above the content again.
TITLE_IN_FILE = False

MAX_SLUG_LENGTH = 50  # leaves room for the unique suffix + extension
MAX_CONTENT_CHARS = 200_000   # ADDED: a runaway model answer cannot build a giant file
MAX_IMAGE_PIXELS = 25_000_000  # CHANGED: was 40 million; each decoded picture can take 4 bytes per pixel of memory

_TABLE_ROW_RE = re.compile(r"^\s*\|(.+)\|\s*$")
_TABLE_SEPARATOR_RE = re.compile(r"^\s*\|?[\s:|-]+\|?\s*$")
_HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
_IMAGE_RE = re.compile(r"^\s*\[\[IMAGE\s+(\d+)\]\]\s*$", re.IGNORECASE)  # where a picture goes
_CONTROL_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
_FORMULA_RE = re.compile(r"^[=+\-@\t\r]")
_NUMBER_RE = re.compile(r"^[+-]?\d+([.,]\d+)*%?$")

# ADDED: generated files were never deleted, so the disk filled up slowly and every old link
# worked forever. Files older than FILE_RETENTION_DAYS (default 30) are removed. After that, the
# download button on a very old conversation shows "not found". Change it in .env, e.g.
# FILE_RETENTION_DAYS=14
_RETENTION_DAYS = float(os.environ.get("FILE_RETENTION_DAYS", "30"))
_last_cleanup = 0.0


def cleanup_old_files() -> int:
    """Deletes generated files past the retention time. Runs at most once an hour. Never raises."""
    global _last_cleanup
    now = time.time()
    if now - _last_cleanup < 3600:
        return 0
    _last_cleanup = now
    removed = 0
    try:
        cutoff = now - _RETENTION_DAYS * 86400
        for name in os.listdir(GENERATED_FILES_DIR):
            if not name.endswith((".docx", ".pdf", ".xlsx", ".csv")):
                continue
            path = os.path.join(GENERATED_FILES_DIR, name)
            try:
                if os.path.isfile(path) and os.path.getmtime(path) < cutoff:
                    os.remove(path)
                    removed += 1
            except OSError:
                pass
    except Exception as e:
        print(f"[file_cleanup] failed: {e}")
    if removed:
        print(f"[file_cleanup] removed {removed} old file(s)")
    return removed


def _clean(text: str) -> str:
    """ADDED: control characters make python-docx / openpyxl raise, so the whole file would fail."""
    return _CONTROL_RE.sub("", text or "")


def _slugify(title: str) -> str:
    """
    Turns arbitrary user-request text into a filesystem-safe name fragment.
    Windows forbids \\ / : * ? " < > | in filenames - a raw title containing
    any of these would crash on save.
    """
    slug = re.sub(r"[^a-z0-9]+", "-", title.lower()).strip("-")
    return slug[:MAX_SLUG_LENGTH] or "document"


def _build_file_name(title: str, fmt: str) -> str:
    """
    Readable-but-collision-safe file name. CHANGED: the random part is 24 characters, not 8.
    /files/ is a public link, so the random part is the only thing stopping a stranger
    from guessing someone else's file name.
    """
    slug = _slugify(title)
    suffix = uuid.uuid4().hex[:24]
    return f"{slug}-{suffix}.{fmt}"


def _split_table_row(line: str) -> list[str]:
    inner = line.strip().strip("|")
    return [cell.strip() for cell in inner.split("|")]


def _safe_cell(text: str) -> str:
    """
    ADDED (CSV): a cell starting with = + - @ is run as a formula when the file is opened in
    Excel. A model answer can contain one (for example =HYPERLINK(...)). A leading ' makes it text.
    Plain numbers like -5 are left alone.
    """
    if _FORMULA_RE.match(text) and not _NUMBER_RE.match(text):
        return "'" + text
    return text


def _normalize_image(data: bytes) -> bytes | None:
    """
    Re-saves any picture (jpg, png, webp, gif...) as a PNG of sensible size, so Word, PDF and
    Excel can all use it. Needs Pillow. Returns None if the picture cannot be read.
    ADDED: the size is checked from the header BEFORE the picture is decoded, so a tiny file
    that expands to billions of pixels cannot use up the server's memory.
    """
    try:
        from PIL import Image
        img = Image.open(io.BytesIO(data))
        if img.width * img.height > MAX_IMAGE_PIXELS:
            print(f"[file_generation] picture too large ({img.width}x{img.height}), skipped")
            return None
        if img.mode not in ("RGB", "RGBA"):
            img = img.convert("RGB")
        img.thumbnail((1600, 1600))
        out = io.BytesIO()
        img.save(out, format="PNG")
        return out.getvalue()
    except Exception as e:
        print(f"[file_generation] could not use a picture: {type(e).__name__}: {e}")
        return None


def _image_at(images: list[bytes] | None, index: int) -> bytes | None:
    if images and 0 <= index < len(images):
        return images[index]
    return None


def parse_structured_content(content: str, image_count: int = 0) -> list[dict]:
    """
    Turns model output into a list of blocks:
      {"type": "heading", "level": int, "text": str}
      {"type": "table", "rows": list[list[str]]}   # rows[0] is the header row
      {"type": "paragraph", "text": str}            # one or more lines, joined
      {"type": "image", "index": int}               # a line like [[IMAGE 1]]
    Table detection: a '|...|' line followed by a separator row (---/:---/---:) is a real table
    and the separator row is dropped. A lone '|...|' line with no separator is a normal paragraph.
    CHANGED: every table row is padded or cut to the header's width. Before, one row with an extra
    '|' made python-docx raise IndexError and the whole file failed.
    Any attached picture the model did not place with [[IMAGE n]] is added at the end.
    """
    lines = content.split("\n")
    blocks = []
    paragraph_buffer = []
    i = 0

    def _flush_paragraph():
        if paragraph_buffer:
            text = "\n".join(paragraph_buffer).strip()
            if text:
                blocks.append({"type": "paragraph", "text": text})
            paragraph_buffer.clear()

    def _fit(row: list[str], width: int) -> list[str]:
        return (row + [""] * width)[:width]

    while i < len(lines):
        line = lines[i]

        image_match = _IMAGE_RE.match(line)
        if image_match:
            _flush_paragraph()
            blocks.append({"type": "image", "index": int(image_match.group(1)) - 1})
            i += 1
            continue

        heading_match = _HEADING_RE.match(line)
        if heading_match:
            _flush_paragraph()
            blocks.append({
                "type": "heading",
                "level": len(heading_match.group(1)),
                "text": heading_match.group(2).strip(),
            })
            i += 1
            continue

        if _TABLE_ROW_RE.match(line) and i + 1 < len(lines) and _TABLE_SEPARATOR_RE.match(lines[i + 1]):
            _flush_paragraph()
            header = _split_table_row(line)
            width = len(header)
            rows = [header]
            i += 2  # skip the header row and the separator row
            while i < len(lines) and _TABLE_ROW_RE.match(lines[i]):
                rows.append(_fit(_split_table_row(lines[i]), width))
                i += 1
            blocks.append({"type": "table", "rows": rows})
            continue

        if line.strip():
            paragraph_buffer.append(line)
        else:
            _flush_paragraph()
        i += 1

    _flush_paragraph()

    placed = {b["index"] for b in blocks if b["type"] == "image"}
    for n in range(image_count):
        if n not in placed:
            blocks.append({"type": "image", "index": n})
    return blocks


# ── ADDED: nicer document formatting ────────────────────────────────────────────────────────
# Before, a model answer with "**bold**", "- bullets" or "1. steps" showed those symbols literally in
# the file, rules like "---" were printed, and the PDF printed black squares for characters such as
# the non-breaking hyphen. parse_for() cleans all of that; Word and PDF get real bold, italics and
# lists, Excel and CSV get plain text.
_HR_RE = re.compile(r"^\s*([-*_])(\s*\1){2,}\s*$")
_BULLET_RE = re.compile(r"^\s*[-*•●]\s+(.*)$")
_NUMBER_ITEM_RE = re.compile(r"^\s*\d{1,3}[.)]\s+(.*)$")
_BOLD_RE = re.compile(r"\*\*(.+?)\*\*")
_ITALIC_RE = re.compile(r"(?<![\w*])\*(?!\s)(.+?)(?<!\s)\*(?![\w*])")
_PDF_MAP = {
    "‐": "-", "‑": "-", "‒": "-", "−": "-", "→": "->", "←": "<-",
    "✓": "", "✔": "", "​": "", "‌": "", "‍": "", "﻿": "",
    " ": " ", " ": " ", " ": " ", " ": " ",
}


def _plain(text: str) -> str:
    """Removes the markdown symbols from text (used for headings, table cells, Excel and CSV)."""
    text = _BOLD_RE.sub(r"\1", text)
    text = _ITALIC_RE.sub(r"\1", text)
    return text.replace("`", "")


def _pdf_safe(text: str) -> str:
    """The standard PDF fonts only hold Western characters. Swap lookalikes, drop the rest
    (emoji, symbols) instead of printing black squares."""
    for a, b in _PDF_MAP.items():
        text = text.replace(a, b)
    return text.encode("cp1252", "ignore").decode("cp1252")


def parse_for(content: str, image_count: int = 0, rich: bool = False) -> list[dict]:
    """
    parse_structured_content() plus formatting clean-up.
    rich=True  (Word, PDF): bullets and numbered steps become {"type": "list", "ordered": bool,
               "items": [...]} blocks; paragraphs keep **bold** and *italic* for the writer to render.
    rich=False (Excel, CSV): everything is plain text; bullets become '• item'.
    Headings and table cells are always plain. Lines like '---' and a leading '> ' are removed.
    """
    lines = []
    for ln in content.split("\n"):
        if _HR_RE.match(ln) and not _TABLE_ROW_RE.match(ln):
            lines.append("")
            continue
        lines.append(re.sub(r"^\s*>\s?", "", ln))
    blocks = parse_structured_content("\n".join(lines), image_count)
    out = []
    for b in blocks:
        if b["type"] == "heading":
            b["text"] = _plain(b["text"])
            out.append(b)
        elif b["type"] == "table":
            b["rows"] = [[_plain(c) for c in row] for row in b["rows"]]
            out.append(b)
        elif b["type"] == "paragraph":
            group, kind = [], None

            def _emit():
                if not group:
                    return
                if kind == "text":
                    t = "\n".join(group)
                    out.append({"type": "paragraph", "text": t if rich else _plain(t)})
                elif rich:
                    out.append({"type": "list", "ordered": kind == "num", "items": list(group)})
                else:
                    for n, item in enumerate(group, 1):
                        out.append({"type": "paragraph",
                                    "text": ("%d. %s" % (n, _plain(item))) if kind == "num" else "• " + _plain(item)})

            for ln in b["text"].split("\n"):
                m_b, m_n = _BULLET_RE.match(ln), _NUMBER_ITEM_RE.match(ln)
                k = "bullet" if m_b else "num" if m_n else "text"
                item = (m_b or m_n).group(1).strip() if (m_b or m_n) else ln
                if k != kind:
                    _emit()
                    group, kind = [], k
                group.append(item)
            _emit()
        else:
            out.append(b)
    return out


def _docx_runs(par, text: str) -> None:
    """Adds text to a Word paragraph with real bold and italic runs."""
    pos = 0
    for m in re.finditer(r"\*\*(.+?)\*\*|(?<![\w*])\*(?!\s)(.+?)(?<!\s)\*(?![\w*])", text):
        if m.start() > pos:
            par.add_run(text[pos:m.start()].replace("`", ""))
        if m.group(1) is not None:
            par.add_run(m.group(1)).bold = True
        else:
            par.add_run(m.group(2)).italic = True
        pos = m.end()
    if pos < len(text):
        par.add_run(text[pos:].replace("`", ""))


def _pdf_inline(text: str) -> str:
    """Escapes text for the PDF writer and turns **bold** / *italic* into real bold and italic."""
    text = _xml_escape(text.replace("`", ""))
    text = _BOLD_RE.sub(r"<b>\1</b>", text)
    text = _ITALIC_RE.sub(r"<i>\1</i>", text)
    return text.replace("\n", "<br/>")


def generate_docx(content: str, title: str, file_name: str, images: list[bytes] | None = None) -> None:
    from docx import Document  # python-docx
    from docx.shared import Inches

    doc = Document()
    doc.core_properties.title = title[:200]  # ADDED: the short title lives in the file's properties, not on the page
    if TITLE_IN_FILE:  # CHANGED: the title is no longer printed on top of the content
        doc.add_heading(title, level=1)
    try:  # ADDED: readable body text: 11 pt, space between paragraphs
        from docx.shared import Pt
        normal = doc.styles["Normal"]
        normal.font.name = "Calibri"
        normal.font.size = Pt(11)
        normal.paragraph_format.space_after = Pt(8)
        normal.paragraph_format.line_spacing = 1.15
    except Exception:
        pass

    for block in parse_for(content, len(images or []), rich=True):
        if block["type"] == "heading":
            doc.add_heading(block["text"], level=min(block["level"] + 1, 9))
        elif block["type"] == "table":
            rows = block["rows"]
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Light Grid Accent 1"
            for r, row_data in enumerate(rows):
                for c, cell_text in enumerate(row_data):
                    cell = table.cell(r, c)
                    cell.text = cell_text
                    if r == 0:
                        for p in cell.paragraphs:
                            for run in p.runs:
                                run.bold = True
        elif block["type"] == "list":
            for item in block["items"]:
                _docx_runs(doc.add_paragraph(style="List Number" if block["ordered"] else "List Bullet"), item)
        elif block["type"] == "image":
            data = _image_at(images, block["index"])
            if data:
                doc.add_picture(io.BytesIO(data), width=Inches(5.5))
        else:
            _docx_runs(doc.add_paragraph(), block["text"])  # CHANGED: real bold and italic

    doc.save(os.path.join(GENERATED_FILES_DIR, file_name))


def generate_pdf(content: str, title: str, file_name: str, images: list[bytes] | None = None) -> None:
    from reportlab.lib.pagesizes import letter  # reportlab
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.platypus import Image as RLImage
    from reportlab.lib.utils import ImageReader
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib import colors

    file_path = os.path.join(GENERATED_FILES_DIR, file_name)
    styles = getSampleStyleSheet()
    # CHANGED: Paragraph reads its text as markup, so a '<' or '&' in the answer (for example
    # "a < b", "R&D", "<div>") raised an error and no PDF was made. Everything is escaped now.
    # CHANGED: the title is no longer printed on top of the content (a letter must start with its own text)
    story = [Paragraph(_xml_escape(title), styles["Title"]), Spacer(1, 12)] if TITLE_IN_FILE else []

    heading_styles = {1: "Heading1", 2: "Heading2", 3: "Heading3"}
    content, title = _pdf_safe(content), _pdf_safe(title)  # ADDED: no black squares for unsupported characters
    from reportlab.lib.styles import ParagraphStyle
    styles["Normal"].fontSize = 11   # ADDED: readable body text
    styles["Normal"].leading = 15
    item_style = ParagraphStyle("ListItem", parent=styles["Normal"], leftIndent=20, bulletIndent=6, spaceAfter=3)

    for block in parse_for(content, len(images or []), rich=True):
        if block["type"] == "heading":
            style_name = heading_styles.get(block["level"], "Heading4")
            story.append(Paragraph(_xml_escape(block["text"]), styles[style_name]))
            story.append(Spacer(1, 6))
        elif block["type"] == "table":
            # CHANGED: cells are Paragraphs so long text wraps instead of running off the page
            rows = [
                [Paragraph(("<b>%s</b>" if r == 0 else "%s") % _xml_escape(cell), styles["Normal"]) for cell in row]
                for r, row in enumerate(block["rows"])
            ]
            table = Table(rows, repeatRows=1)
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e0e0e0")),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]))
            story.append(table)
            story.append(Spacer(1, 12))
        elif block["type"] == "list":
            for n, item in enumerate(block["items"], 1):
                story.append(Paragraph(_pdf_inline(item), item_style, bulletText=("%d." % n) if block["ordered"] else "•"))
            story.append(Spacer(1, 6))
        elif block["type"] == "image":
            data = _image_at(images, block["index"])
            if data:
                iw, ih = ImageReader(io.BytesIO(data)).getSize()
                scale = min(450 / iw, 600 / ih, 1)  # keep it inside the page margins
                story.append(RLImage(io.BytesIO(data), width=iw * scale, height=ih * scale))
                story.append(Spacer(1, 12))
        else:
            story.append(Paragraph(_pdf_inline(block["text"]), styles["Normal"]))  # CHANGED: real bold and italic
            story.append(Spacer(1, 8))

    doc = SimpleDocTemplate(file_path, pagesize=letter)
    doc.build(story)


def _put(ws, row: int, col: int, value: str):
    """ADDED (XLSX): openpyxl turns any text starting with '=' into a live formula. Force plain text."""
    cell = ws.cell(row=row, column=col, value=value)
    if isinstance(value, str):
        cell.data_type = "s"
    return cell


def generate_xlsx(content: str, title: str, file_name: str, images: list[bytes] | None = None) -> None:
    from openpyxl import Workbook  # openpyxl
    from openpyxl.styles import Font
    from openpyxl.drawing.image import Image as XLImage

    wb = Workbook()
    ws = wb.active
    # CHANGED: Excel forbids  [ ] : * ? / \  in a sheet name, and openpyxl raises if one is present
    ws.title = re.sub(r"[\[\]:*?/\\]", " ", title).strip(" '")[:31] or "Sheet1"

    row_num = 1
    for block in parse_for(content, len(images or []), rich=False):
        if block["type"] == "heading":
            cell = _put(ws, row_num, 1, block["text"])
            cell.font = Font(bold=True, size=14 - block["level"])
            row_num += 1
        elif block["type"] == "table":
            for r, row_data in enumerate(block["rows"]):
                for c, cell_text in enumerate(row_data, start=1):
                    cell = _put(ws, row_num, c, cell_text)
                    if r == 0:
                        cell.font = Font(bold=True)
                row_num += 1
        elif block["type"] == "image":
            data = _image_at(images, block["index"])
            if data:
                picture = XLImage(io.BytesIO(data))
                scale = min(1, 480 / picture.width)
                picture.width, picture.height = picture.width * scale, picture.height * scale
                ws.add_image(picture, f"A{row_num}")
                row_num += int(picture.height / 20) + 1  # leave room so rows below do not sit under the picture
        else:
            for line in block["text"].split("\n"):
                if line.strip():
                    _put(ws, row_num, 1, line)
                    row_num += 1
        row_num += 1  # blank row between blocks

    wb.save(os.path.join(GENERATED_FILES_DIR, file_name))


def generate_csv(content: str, title: str, file_name: str, images: list[bytes] | None = None) -> None:
    file_path = os.path.join(GENERATED_FILES_DIR, file_name)
    # utf-8-sig: Excel on Windows shows accents/Filipino characters correctly only with the BOM
    with open(file_path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv_module.writer(f)
        for block in parse_for(content, 0, rich=False):  # a CSV file cannot hold pictures
            if block["type"] == "heading":
                writer.writerow([_safe_cell(block["text"])])
            elif block["type"] == "table":
                for row_data in block["rows"]:
                    writer.writerow([_safe_cell(c) for c in row_data])
            elif block["type"] == "paragraph":  # CHANGED: a [[IMAGE n]] line used to crash CSV with KeyError
                for line in block["text"].split("\n"):
                    if line.strip():
                        writer.writerow([_safe_cell(line)])


_FILE_GENERATORS = {
    "docx": generate_docx,
    "pdf": generate_pdf,
    "xlsx": generate_xlsx,
    "csv": generate_csv,
}


def generate_file(content: str, fmt: str, title: str = "Document",
                  images: list[bytes] | None = None) -> tuple[str, str]:
    """
    Dispatches to the right format-specific generator. The file name is
    built once, here, from the user's actual request (via title).
    images: raw bytes of pictures the user attached; they are placed where the
    content says [[IMAGE 1]], [[IMAGE 2]]... or at the end if not placed.
    Returns (public_url_path, file_name).
    """
    cleanup_old_files()
    generator = _FILE_GENERATORS.get(fmt)
    if generator is None:
        raise ValueError(f"Unsupported file format: {fmt!r}")

    content = _clean(content)[:MAX_CONTENT_CHARS]
    title = _clean(title)[:200] or "Document"
    usable = [p for p in (_normalize_image(raw) for raw in (images or [])) if p]

    file_name = _build_file_name(title, fmt)
    generator(content, title, file_name, usable)
=======
import csv as csv_module
import io
import os
import re
import uuid

GENERATED_FILES_DIR = "generated_files"
os.makedirs(GENERATED_FILES_DIR, exist_ok=True)

MAX_SLUG_LENGTH = 50  # leaves room for the unique suffix + extension

_TABLE_ROW_RE = re.compile(r"^\s*\|(.+)\|\s*$")
_TABLE_SEPARATOR_RE = re.compile(r"^\s*\|?[\s:|-]+\|?\s*$")
_HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
_IMAGE_RE = re.compile(r"^\s*\[\[IMAGE\s+(\d+)\]\]\s*$", re.IGNORECASE)  # ADDED: where a picture goes


def _slugify(title: str) -> str:
    """
    Turns arbitrary user-request text into a filesystem-safe name fragment.
    Windows forbids \\ / : * ? " < > | in filenames - a raw title containing
    any of these would crash on save.
    """
    slug = re.sub(r"[^a-z0-9]+", "-", title.lower()).strip("-")
    return slug[:MAX_SLUG_LENGTH] or "document"


def _build_file_name(title: str, fmt: str) -> str:
    """
    Readable-but-collision-safe file name: a slug of the user's actual
    request, plus a short random suffix, so two similar requests never
    silently overwrite each other's files.
    """
    slug = _slugify(title)
    suffix = uuid.uuid4().hex[:8]
    return f"{slug}-{suffix}.{fmt}"


def _split_table_row(line: str) -> list[str]:
    inner = line.strip().strip("|")
    return [cell.strip() for cell in inner.split("|")]


def _normalize_image(data: bytes) -> bytes | None:
    """
    ADDED: re-saves any picture (jpg, png, webp, gif...) as a PNG of sensible
    size, so Word, PDF and Excel can all use it. Needs Pillow (pip install pillow).
    Returns None if the picture cannot be read.
    """
    try:
        from PIL import Image
        img = Image.open(io.BytesIO(data))
        if img.mode not in ("RGB", "RGBA"):
            img = img.convert("RGB")
        img.thumbnail((1600, 1600))
        out = io.BytesIO()
        img.save(out, format="PNG")
        return out.getvalue()
    except Exception as e:
        print(f"[file_generation] could not use a picture: {type(e).__name__}: {e}")
        return None


def _image_at(images: list[bytes] | None, index: int) -> bytes | None:
    if images and 0 <= index < len(images):
        return images[index]
    return None


def parse_structured_content(content: str, image_count: int = 0) -> list[dict]:
    """
    Turns model output into a list of blocks:
      {"type": "heading", "level": int, "text": str}
      {"type": "table", "rows": list[list[str]]}   # rows[0] is the header row
      {"type": "paragraph", "text": str}            # one or more lines, joined
      {"type": "image", "index": int}               # ADDED: a line like [[IMAGE 1]]
    Table detection: two or more consecutive '|...|' lines, where the second
    line is a separator row (---/:---/---:) gets treated as a real table and
    that separator row is dropped from the output rows. A lone '|...|' line
    with no separator underneath it is NOT treated as a table - falls through
    to a normal paragraph, since that's more likely stray text than a real
    table the model intended.
    Any attached picture the model did not place with [[IMAGE n]] is added at the end.
    """
    lines = content.split("\n")
    blocks = []
    paragraph_buffer = []
    i = 0

    def _flush_paragraph():
        if paragraph_buffer:
            text = "\n".join(paragraph_buffer).strip()
            if text:
                blocks.append({"type": "paragraph", "text": text})
            paragraph_buffer.clear()

    while i < len(lines):
        line = lines[i]

        image_match = _IMAGE_RE.match(line)
        if image_match:
            _flush_paragraph()
            blocks.append({"type": "image", "index": int(image_match.group(1)) - 1})
            i += 1
            continue

        heading_match = _HEADING_RE.match(line)
        if heading_match:
            _flush_paragraph()
            blocks.append({
                "type": "heading",
                "level": len(heading_match.group(1)),
                "text": heading_match.group(2).strip(),
            })
            i += 1
            continue

        if _TABLE_ROW_RE.match(line) and i + 1 < len(lines) and _TABLE_SEPARATOR_RE.match(lines[i + 1]):
            _flush_paragraph()
            rows = [_split_table_row(line)]
            i += 2  # skip the header row and the separator row
            while i < len(lines) and _TABLE_ROW_RE.match(lines[i]):
                rows.append(_split_table_row(lines[i]))
                i += 1
            blocks.append({"type": "table", "rows": rows})
            continue

        if line.strip():
            paragraph_buffer.append(line)
        else:
            _flush_paragraph()
        i += 1

    _flush_paragraph()

    placed = {b["index"] for b in blocks if b["type"] == "image"}
    for n in range(image_count):
        if n not in placed:
            blocks.append({"type": "image", "index": n})
    return blocks


def generate_docx(content: str, title: str, file_name: str, images: list[bytes] | None = None) -> None:
    from docx import Document  # python-docx
    from docx.shared import Inches

    doc = Document()
    doc.add_heading(title, level=1)

    for block in parse_structured_content(content, len(images or [])):
        if block["type"] == "heading":
            doc.add_heading(block["text"], level=min(block["level"] + 1, 9))
        elif block["type"] == "table":
            rows = block["rows"]
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Light Grid Accent 1"
            for r, row_data in enumerate(rows):
                for c, cell_text in enumerate(row_data):
                    cell = table.cell(r, c)
                    cell.text = cell_text
                    if r == 0:
                        for p in cell.paragraphs:
                            for run in p.runs:
                                run.bold = True
        elif block["type"] == "image":
            data = _image_at(images, block["index"])
            if data:
                doc.add_picture(io.BytesIO(data), width=Inches(5.5))
        else:
            doc.add_paragraph(block["text"])

    doc.save(os.path.join(GENERATED_FILES_DIR, file_name))


def generate_pdf(content: str, title: str, file_name: str, images: list[bytes] | None = None) -> None:
    from reportlab.lib.pagesizes import letter  # reportlab
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.platypus import Image as RLImage
    from reportlab.lib.utils import ImageReader
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib import colors

    file_path = os.path.join(GENERATED_FILES_DIR, file_name)
    styles = getSampleStyleSheet()
    story = [Paragraph(title, styles["Title"]), Spacer(1, 12)]

    heading_styles = {1: "Heading1", 2: "Heading2", 3: "Heading3"}

    for block in parse_structured_content(content, len(images or [])):
        if block["type"] == "heading":
            style_name = heading_styles.get(block["level"], "Heading4")
            story.append(Paragraph(block["text"], styles[style_name]))
            story.append(Spacer(1, 6))
        elif block["type"] == "table":
            table = Table(block["rows"])
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e0e0e0")),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]))
            story.append(table)
            story.append(Spacer(1, 12))
        elif block["type"] == "image":
            data = _image_at(images, block["index"])
            if data:
                iw, ih = ImageReader(io.BytesIO(data)).getSize()
                scale = min(450 / iw, 600 / ih, 1)  # keep it inside the page margins
                story.append(RLImage(io.BytesIO(data), width=iw * scale, height=ih * scale))
                story.append(Spacer(1, 12))
        else:
            story.append(Paragraph(block["text"].replace("\n", "<br/>"), styles["Normal"]))
            story.append(Spacer(1, 6))

    doc = SimpleDocTemplate(file_path, pagesize=letter)
    doc.build(story)


def generate_xlsx(content: str, title: str, file_name: str, images: list[bytes] | None = None) -> None:
    from openpyxl import Workbook  # openpyxl
    from openpyxl.styles import Font
    from openpyxl.drawing.image import Image as XLImage

    wb = Workbook()
    ws = wb.active
    ws.title = title[:31] or "Sheet1"  # Excel sheet names cap at 31 chars

    row_num = 1
    for block in parse_structured_content(content, len(images or [])):
        if block["type"] == "heading":
            cell = ws.cell(row=row_num, column=1, value=block["text"])
            cell.font = Font(bold=True, size=14 - block["level"])
            row_num += 1
        elif block["type"] == "table":
            for r, row_data in enumerate(block["rows"]):
                for c, cell_text in enumerate(row_data, start=1):
                    cell = ws.cell(row=row_num, column=c, value=cell_text)
                    if r == 0:
                        cell.font = Font(bold=True)
                row_num += 1
        elif block["type"] == "image":
            data = _image_at(images, block["index"])
            if data:
                picture = XLImage(io.BytesIO(data))
                scale = min(1, 480 / picture.width)
                picture.width, picture.height = picture.width * scale, picture.height * scale
                ws.add_image(picture, f"A{row_num}")
                row_num += int(picture.height / 20) + 1  # leave room so rows below do not sit under the picture
        else:
            for line in block["text"].split("\n"):
                if line.strip():
                    ws.cell(row=row_num, column=1, value=line)
                    row_num += 1
        row_num += 1  # blank row between blocks

    wb.save(os.path.join(GENERATED_FILES_DIR, file_name))


def generate_csv(content: str, title: str, file_name: str, images: list[bytes] | None = None) -> None:
    file_path = os.path.join(GENERATED_FILES_DIR, file_name)
    with open(file_path, "w", newline="", encoding="utf-8") as f:
        writer = csv_module.writer(f)
        for block in parse_structured_content(content):  # a CSV file cannot hold pictures
            if block["type"] == "heading":
                writer.writerow([block["text"]])
            elif block["type"] == "table":
                for row_data in block["rows"]:
                    writer.writerow(row_data)
            else:
                for line in block["text"].split("\n"):
                    if line.strip():
                        writer.writerow([line])


_FILE_GENERATORS = {
    "docx": generate_docx,
    "pdf": generate_pdf,
    "xlsx": generate_xlsx,
    "csv": generate_csv,
}


def generate_file(content: str, fmt: str, title: str = "Document",
                  images: list[bytes] | None = None) -> tuple[str, str]:
    """
    Dispatches to the right format-specific generator. The file name is
    built once, here, from the user's actual request (via title).
    images: raw bytes of pictures the user attached; they are placed where the
    content says [[IMAGE 1]], [[IMAGE 2]]... or at the end if not placed.
    Returns (public_url_path, file_name).
    """
    generator = _FILE_GENERATORS.get(fmt)
    if generator is None:
        raise ValueError(f"Unsupported file format: {fmt!r}")

    usable = [p for p in (_normalize_image(raw) for raw in (images or [])) if p]

    file_name = _build_file_name(title, fmt)
    generator(content, title, file_name, usable)
>>>>>>> 17d14f619ee27a032db336356d927b1b42009a91
    return f"/files/{file_name}", file_name