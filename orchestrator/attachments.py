<<<<<<< HEAD
"""
Reads files and images the user attaches to a chat message.
- Documents: text is extracted on the server.
- Images: a vision model OBSERVES and DESCRIBES them. Nothing is ever generated or edited.
"""

import base64
import io
import os
import re
import uuid
import zipfile

import requests

# Copy the current model name from Google's "Image understanding" docs for the
# same Gemini API key you already use for embeddings. Names change over time.
VISION_MODEL = "gemini-3-flash-preview"

MAX_FILE_CHARS = 12000  # keeps token use (and your 50,000 token quota) under control
MAX_PDF_PAGES = 50
MAX_UNZIPPED_BYTES = 50 * 1024 * 1024  # a small .docx/.xlsx can be a zip bomb that unpacks to gigabytes
ALLOWED_EXT = (".txt", ".md", ".csv", ".pdf", ".docx", ".xlsx")
ALLOWED_IMAGES = ("image/png", "image/jpeg", "image/webp", "image/gif", "image/heic", "image/heif")


def check_attachments(attachments: list[dict]) -> str | None:
    """Returns a short error text for the first file that is not allowed, otherwise None."""
    for a in attachments:
        name = str(a.get("name") or "file")
        mime = str(a.get("type") or "")
        ok_image = mime in ALLOWED_IMAGES
        ok_doc = name.lower().endswith(ALLOWED_EXT) and not mime.startswith("image/")
        if not (ok_image or ok_doc):
            return f"{name[:60]}: this file type is not supported"
    return None


_DOC_MIME = {
    ".txt": "text/plain", ".md": "text/plain", ".csv": "text/csv", ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}


def _stored_mime(a: dict) -> str:
    """The type a file is stored with. Never the browser's claim for documents, so nobody can store a page that runs as HTML."""
    mime = str(a.get("type") or "")
    if mime in ALLOWED_IMAGES:
        return mime
    name = str(a.get("name") or "").lower()
    return next((m for ext, m in _DOC_MIME.items() if name.endswith(ext)), "application/octet-stream")


def _zip_too_big(data: bytes) -> bool:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            return sum(i.file_size for i in z.infolist()) > MAX_UNZIPPED_BYTES
    except Exception:
        return True


def extract_text(name: str, data: bytes) -> str:
    n = name.lower()

    if n.endswith((".txt", ".md", ".csv")):
        return data[: MAX_FILE_CHARS * 4].decode("utf-8", errors="ignore")

    if n.endswith(".pdf"):
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        out, size = [], 0
        for i, page in enumerate(reader.pages):
            if i >= MAX_PDF_PAGES or size >= MAX_FILE_CHARS:
                break
            text = page.extract_text() or ""
            out.append(text)
            size += len(text)
        return "\n".join(out)

    if n.endswith(".docx"):
        from docx import Document
        if _zip_too_big(data):
            return ""
        doc = Document(io.BytesIO(data))
        lines, size = [], 0
        for p in doc.paragraphs:
            lines.append(p.text)
            size += len(p.text)
            if size >= MAX_FILE_CHARS:
                return "\n".join(lines)
        for table in doc.tables:
            for row in table.rows:
                line = " | ".join(cell.text for cell in row.cells)
                lines.append(line)
                size += len(line)
                if size >= MAX_FILE_CHARS:
                    return "\n".join(lines)
        return "\n".join(lines)

    if n.endswith(".xlsx"):
        from openpyxl import load_workbook
        if _zip_too_big(data):
            return ""
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        lines, size = [], 0
        for ws in wb.worksheets:
            lines.append(f"[Sheet: {ws.title}]")
            for row in ws.iter_rows(values_only=True):
                line = " | ".join("" if c is None else str(c) for c in row)
                lines.append(line)
                size += len(line)
                if size >= MAX_FILE_CHARS:
                    return "\n".join(lines)
        return "\n".join(lines)

    return ""


def observe_image(b64_data: str, mime: str) -> str | None:
    api_key = os.environ.get("GEMINI_API_KEY")
    if not api_key:
        print("[attachments] GEMINI_API_KEY not set, cannot read image")
        return None

    url = (
        "https://generativelanguage.googleapis.com/v1beta/models/"
        f"{VISION_MODEL}:generateContent"
    )
    body = {
        "contents": [{
            "parts": [
                {"text": (
                    "Observe this image and describe everything useful in it: "
                    "objects, people, any text (transcribe it exactly), charts, "
                    "numbers, layout, and colors. Only describe what is actually "
                    "visible. Do not create, edit, or generate any image."
                )},
                {"inline_data": {"mime_type": mime, "data": b64_data}},
            ]
        }]
    }
    try:
        response = requests.post(
            url, json=body, headers={"x-goog-api-key": api_key}, timeout=40
        )
        response.raise_for_status()
        return response.json()["candidates"][0]["content"]["parts"][0]["text"]
    except Exception as e:
        # type only: the full exception text can contain request details
        print(f"[attachments] image read failed: {type(e).__name__}")
        return None


def build_attachment_context(attachments: list[dict]) -> str:
    parts = []
    for a in attachments:
        name = re.sub(r'[\r\n"]', " ", str(a.get("name") or "file"))[:120]
        mime = str(a.get("type") or "")
        b64 = a.get("data", "")
        try:
            data = base64.b64decode(b64)
        except Exception:
            continue

        if mime.startswith("image/"):
            description = observe_image(b64, mime)
            parts.append(
                f'[Image "{name}" - what it shows]\n'
                f'{description or "Could not read this image."}'
            )
        else:
            try:
                text = extract_text(name, data)[:MAX_FILE_CHARS]
            except Exception as e:  # a damaged file must not break the whole chat
                print(f"[attachments] could not read a file: {type(e).__name__}")
                text = ""
            parts.append(
                f'[File "{name}" - content]\n{text or "Could not read this file."}'
            )
    return "\n\n".join(parts)


UPLOAD_BUCKET = "user-uploads"


def _build_path(user_id: str, conversation_id: str | None, name: str) -> str:
    safe_name = re.sub(r"[^A-Za-z0-9._-]", "_", name)
    return f"{user_id}/{conversation_id or 'no-conversation'}/{uuid.uuid4().hex}-{safe_name}"


def prepare_uploads(user_id: str, conversation_id: str | None, attachments: list[dict]) -> list[dict]:
    """
    Decides each file's storage path BEFORE saving, so the same path can be
    written into the chat message (to show the file again after a refresh)
    and used for the real upload.
    """
    planned = []
    for a in attachments:
        name = str(a.get("name") or "file")[:120]
        planned.append({
            "name": name,
            "type": _stored_mime(a),
            "path": _build_path(user_id, conversation_id, name),
        })
    return planned


def save_attachments(supabase, user_id: str, conversation_id: str | None,
                     attachments: list[dict], planned: list[dict]) -> None:
    """
    Keeps the ORIGINAL files each user sends: the file goes into the private
    Storage bucket under the user's own folder, and one row per file goes into
    the user_uploads table. Runs in the background, so a failure here never
    blocks the chat reply.
    """
    for a, p in zip(attachments, planned):
        try:
            data = base64.b64decode(a.get("data", ""))
            supabase.storage.from_(UPLOAD_BUCKET).upload(
                p["path"], data, {"content-type": p["type"]}
            )
            supabase.table("user_uploads").insert({
                "user_id": user_id,
                "conversation_id": conversation_id,
                "file_name": p["name"],
                "mime_type": p["type"],
                "storage_path": p["path"],
                "size_bytes": len(data),
            }).execute()
        except Exception as e:
            print(f"[uploads] failed to save {p['name']}: {type(e).__name__}: {e}")


def signed_url(supabase, path: str, seconds: int = 3600) -> str | None:
    """Temporary link to a private file. Expires, so it is made fresh each time a chat is opened."""
    try:
        res = supabase.storage.from_(UPLOAD_BUCKET).create_signed_url(path, seconds)
        url = res.get("signedURL") or res.get("signedUrl")
        if url and url.startswith("/"):
            url = os.environ["SUPABASE_URL"].rstrip("/") + "/storage/v1" + url
        return url
    except Exception as e:
        print(f"[uploads] signed url failed for {path}: {type(e).__name__}")
=======
"""
Reads files and images the user attaches to a chat message.
- Documents: text is extracted on the server.
- Images: a vision model OBSERVES and DESCRIBES them. Nothing is ever generated or edited.
"""

import base64
import io
import os
import re
import uuid
import zipfile

import requests

# Copy the current model name from Google's "Image understanding" docs for the
# same Gemini API key you already use for embeddings. Names change over time.
VISION_MODEL = "gemini-3-flash-preview"

MAX_FILE_CHARS = 12000  # keeps token use (and your 50,000 token quota) under control
MAX_PDF_PAGES = 50
MAX_UNZIPPED_BYTES = 50 * 1024 * 1024  # a small .docx/.xlsx can be a zip bomb that unpacks to gigabytes
ALLOWED_EXT = (".txt", ".md", ".csv", ".pdf", ".docx", ".xlsx")
ALLOWED_IMAGES = ("image/png", "image/jpeg", "image/webp", "image/gif", "image/heic", "image/heif")


def check_attachments(attachments: list[dict]) -> str | None:
    """Returns a short error text for the first file that is not allowed, otherwise None."""
    for a in attachments:
        name = str(a.get("name") or "file")
        mime = str(a.get("type") or "")
        ok_image = mime in ALLOWED_IMAGES
        ok_doc = name.lower().endswith(ALLOWED_EXT) and not mime.startswith("image/")
        if not (ok_image or ok_doc):
            return f"{name[:60]}: this file type is not supported"
    return None


_DOC_MIME = {
    ".txt": "text/plain", ".md": "text/plain", ".csv": "text/csv", ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}


def _stored_mime(a: dict) -> str:
    """The type a file is stored with. Never the browser's claim for documents, so nobody can store a page that runs as HTML."""
    mime = str(a.get("type") or "")
    if mime in ALLOWED_IMAGES:
        return mime
    name = str(a.get("name") or "").lower()
    return next((m for ext, m in _DOC_MIME.items() if name.endswith(ext)), "application/octet-stream")


def _zip_too_big(data: bytes) -> bool:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            return sum(i.file_size for i in z.infolist()) > MAX_UNZIPPED_BYTES
    except Exception:
        return True


def extract_text(name: str, data: bytes) -> str:
    n = name.lower()

    if n.endswith((".txt", ".md", ".csv")):
        return data[: MAX_FILE_CHARS * 4].decode("utf-8", errors="ignore")

    if n.endswith(".pdf"):
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        out, size = [], 0
        for i, page in enumerate(reader.pages):
            if i >= MAX_PDF_PAGES or size >= MAX_FILE_CHARS:
                break
            text = page.extract_text() or ""
            out.append(text)
            size += len(text)
        return "\n".join(out)

    if n.endswith(".docx"):
        from docx import Document
        if _zip_too_big(data):
            return ""
        doc = Document(io.BytesIO(data))
        lines, size = [], 0
        for p in doc.paragraphs:
            lines.append(p.text)
            size += len(p.text)
            if size >= MAX_FILE_CHARS:
                return "\n".join(lines)
        for table in doc.tables:
            for row in table.rows:
                line = " | ".join(cell.text for cell in row.cells)
                lines.append(line)
                size += len(line)
                if size >= MAX_FILE_CHARS:
                    return "\n".join(lines)
        return "\n".join(lines)

    if n.endswith(".xlsx"):
        from openpyxl import load_workbook
        if _zip_too_big(data):
            return ""
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        lines, size = [], 0
        for ws in wb.worksheets:
            lines.append(f"[Sheet: {ws.title}]")
            for row in ws.iter_rows(values_only=True):
                line = " | ".join("" if c is None else str(c) for c in row)
                lines.append(line)
                size += len(line)
                if size >= MAX_FILE_CHARS:
                    return "\n".join(lines)
        return "\n".join(lines)

    return ""


def observe_image(b64_data: str, mime: str) -> str | None:
    api_key = os.environ.get("GEMINI_API_KEY")
    if not api_key:
        print("[attachments] GEMINI_API_KEY not set, cannot read image")
        return None

    url = (
        "https://generativelanguage.googleapis.com/v1beta/models/"
        f"{VISION_MODEL}:generateContent"
    )
    body = {
        "contents": [{
            "parts": [
                {"text": (
                    "Observe this image and describe everything useful in it: "
                    "objects, people, any text (transcribe it exactly), charts, "
                    "numbers, layout, and colors. Only describe what is actually "
                    "visible. Do not create, edit, or generate any image."
                )},
                {"inline_data": {"mime_type": mime, "data": b64_data}},
            ]
        }]
    }
    try:
        response = requests.post(
            url, json=body, headers={"x-goog-api-key": api_key}, timeout=40
        )
        response.raise_for_status()
        return response.json()["candidates"][0]["content"]["parts"][0]["text"]
    except Exception as e:
        # type only: the full exception text can contain request details
        print(f"[attachments] image read failed: {type(e).__name__}")
        return None


def build_attachment_context(attachments: list[dict]) -> str:
    parts = []
    for a in attachments:
        name = re.sub(r'[\r\n"]', " ", str(a.get("name") or "file"))[:120]
        mime = str(a.get("type") or "")
        b64 = a.get("data", "")
        try:
            data = base64.b64decode(b64)
        except Exception:
            continue

        if mime.startswith("image/"):
            description = observe_image(b64, mime)
            parts.append(
                f'[Image "{name}" - what it shows]\n'
                f'{description or "Could not read this image."}'
            )
        else:
            try:
                text = extract_text(name, data)[:MAX_FILE_CHARS]
            except Exception as e:  # a damaged file must not break the whole chat
                print(f"[attachments] could not read a file: {type(e).__name__}")
                text = ""
            parts.append(
                f'[File "{name}" - content]\n{text or "Could not read this file."}'
            )
    return "\n\n".join(parts)


UPLOAD_BUCKET = "user-uploads"


def _build_path(user_id: str, conversation_id: str | None, name: str) -> str:
    safe_name = re.sub(r"[^A-Za-z0-9._-]", "_", name)
    return f"{user_id}/{conversation_id or 'no-conversation'}/{uuid.uuid4().hex}-{safe_name}"


def prepare_uploads(user_id: str, conversation_id: str | None, attachments: list[dict]) -> list[dict]:
    """
    Decides each file's storage path BEFORE saving, so the same path can be
    written into the chat message (to show the file again after a refresh)
    and used for the real upload.
    """
    planned = []
    for a in attachments:
        name = str(a.get("name") or "file")[:120]
        planned.append({
            "name": name,
            "type": _stored_mime(a),
            "path": _build_path(user_id, conversation_id, name),
        })
    return planned


def save_attachments(supabase, user_id: str, conversation_id: str | None,
                     attachments: list[dict], planned: list[dict]) -> None:
    """
    Keeps the ORIGINAL files each user sends: the file goes into the private
    Storage bucket under the user's own folder, and one row per file goes into
    the user_uploads table. Runs in the background, so a failure here never
    blocks the chat reply.
    """
    for a, p in zip(attachments, planned):
        try:
            data = base64.b64decode(a.get("data", ""))
            supabase.storage.from_(UPLOAD_BUCKET).upload(
                p["path"], data, {"content-type": p["type"]}
            )
            supabase.table("user_uploads").insert({
                "user_id": user_id,
                "conversation_id": conversation_id,
                "file_name": p["name"],
                "mime_type": p["type"],
                "storage_path": p["path"],
                "size_bytes": len(data),
            }).execute()
        except Exception as e:
            print(f"[uploads] failed to save {p['name']}: {type(e).__name__}: {e}")


def signed_url(supabase, path: str, seconds: int = 3600) -> str | None:
    """Temporary link to a private file. Expires, so it is made fresh each time a chat is opened."""
    try:
        res = supabase.storage.from_(UPLOAD_BUCKET).create_signed_url(path, seconds)
        url = res.get("signedURL") or res.get("signedUrl")
        if url and url.startswith("/"):
            url = os.environ["SUPABASE_URL"].rstrip("/") + "/storage/v1" + url
        return url
    except Exception as e:
        print(f"[uploads] signed url failed for {path}: {type(e).__name__}")
>>>>>>> 17d14f619ee27a032db336356d927b1b42009a91
        return None